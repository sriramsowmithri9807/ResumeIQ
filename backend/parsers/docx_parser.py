"""
DOCX text extraction using python-docx.
Extracts paragraphs and table cells for full coverage.
"""
from docx import Document
import io
import re


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract all text from a DOCX file given as bytes.
    Handles paragraphs and tables.
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
        text_parts = []

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Extract text from tables (skills tables, etc.)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    text_parts.append(row_text)

        raw_text = "\n".join(text_parts)
        return clean_text(raw_text)

    except Exception as e:
        raise ValueError(f"Failed to parse DOCX: {str(e)}")


def clean_text(text: str) -> str:
    """Remove excessive whitespace while preserving structure."""
    text = re.sub(r'\n{3,}', '\n\n', text)
    lines = [line.rstrip() for line in text.split('\n')]
    return '\n'.join(lines).strip()
